from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(r"C:\Users\USER\Documents\Project_TB")
OUTPUT = ROOT / "TB_프로젝트_사업계획_종합정리.docx"
TABLE_HELPERS = Path(
    r"C:\Users\USER\.codex\plugins\cache\openai-primary-runtime\documents\26.826.12353"
    r"\skills\documents\scripts"
)
sys.path.insert(0, str(TABLE_HELPERS))
from table_geometry import apply_table_geometry, audit_docx_tables  # noqa: E402


FONT = "Malgun Gothic"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
INK = "20252B"
MUTED = "5B6573"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
PALE_BLUE = "EEF5FB"
PALE_GOLD = "FFF6DD"
PALE_RED = "FDECEC"
PALE_GREEN = "EAF5EE"
WHITE = "FFFFFF"
RED = "9B1C1C"
GOLD = "7A5A00"
GREEN = "2F6B45"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 100, "bottom": 100, "start": 120, "end": 120}


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    if run._element.get_or_add_rPr().rFonts is None:
        run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
    rfonts = run._element.get_or_add_rPr().rFonts
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style(style, size, color=INK, bold=False, before=0, after=6, line_spacing=1.10):
    style.font.name = FONT
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing
    style.paragraph_format.widow_control = True


def set_paragraph_border(paragraph, color="D9E2EA", size="6", space="5"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_text(cell, text, *, bold=False, color=INK, size=9.3, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def add_numbering_definition(doc, *, num_format, level_text):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 0
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    lvl.append(fmt)
    txt = OxmlElement("w:lvlText")
    txt.set(qn("w:val"), level_text)
    lvl.append(txt)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    r_pr.append(r_fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def set_abstract_numbering_indent(doc, abstract_id, *, left=720, hanging=360, tab_pos=720):
    numbering = doc.part.numbering_part.element
    target = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        if abstract.get(qn("w:abstractNumId")) == str(abstract_id):
            target = abstract
            break
    if target is None:
        raise ValueError(f"abstract numbering definition {abstract_id} not found")
    lvl = target.find(qn("w:lvl"))
    p_pr = lvl.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        lvl.append(p_pr)
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        p_pr.append(tabs)
    for child in list(tabs):
        tabs.remove(child)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(tab_pos))
    tabs.append(tab)
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))


def add_numbering_instance(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    num_id = (max(num_ids) + 1) if num_ids else 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override)
    num.append(lvl_override)
    numbering.append(num)
    return num_id


def add_list_paragraph(doc, text, num_id, *, bold_prefix=None, after=8, color=INK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.167
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix) :])
        set_run_font(r2, size=11, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=color)
    return p


def add_body(doc, text, *, bold_prefix=None, italic=False, color=INK, after=6, align="left"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix) :])
        set_run_font(r2, size=11, italic=italic, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, italic=italic, color=color)
    return p


def add_source(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(f"근거: {text}")
    set_run_font(r, size=8.5, italic=True, color=MUTED)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(text, style="Caption")
    p.paragraph_format.keep_with_next = True
    return p


def add_callout(doc, label, text, *, fill=PALE_BLUE, label_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    apply_table_geometry(
        table,
        [CONTENT_WIDTH_DXA],
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=180,
        cell_margins_dxa={"top": 150, "bottom": 150, "start": 180, "end": 180},
    )
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(f"{label}  ")
    set_run_font(r1, size=10.5, bold=True, color=label_color)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)
    return table


def add_table(doc, headers, rows, widths, *, aligns=None, font_size=9.3, shade_first_col=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color=NAVY, size=9.2, align="center")
        set_cell_shading(table.rows[0].cells[idx], LIGHT_GRAY)
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])

    resolved_aligns = aligns or ["left"] * len(headers)
    for row_idx, row_data in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for col_idx, value in enumerate(row_data):
            set_cell_text(cells[col_idx], value, size=font_size, align=resolved_aligns[col_idx])
            if shade_first_col and col_idx == 0:
                set_cell_shading(cells[col_idx], BLUE_GRAY)
                for run in cells[col_idx].paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        if row_idx % 2 == 1 and not shade_first_col:
            for cell in cells:
                set_cell_shading(cell, "FAFBFC")

    apply_table_geometry(
        table,
        widths,
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS,
    )
    table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    props = doc.core_properties
    props.title = "TB 프로젝트 사업계획 종합 정리"
    props.subject = "결핵 진단 AI 모델 공동개발 사업의 실행 기준·리스크·결정사항 정리"
    props.author = "Project TB Working Brief"
    props.keywords = "결핵, AI, 사업계획, 대한결핵협회, 마이허브, 인허가"

    set_style(doc.styles["Normal"], 11, INK, False, 0, 6, 1.10)
    set_style(doc.styles["Heading 1"], 16, BLUE, True, 16, 8, 1.05)
    set_style(doc.styles["Heading 2"], 13, BLUE, True, 12, 6, 1.05)
    set_style(doc.styles["Heading 3"], 12, DARK_BLUE, True, 8, 4, 1.05)
    set_style(doc.styles["Caption"], 9, MUTED, True, 4, 4, 1.0)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.styles["Heading 2"].paragraph_format.keep_with_next = True
    doc.styles["Heading 3"].paragraph_format.keep_with_next = True

    # Use Word's native single-level List Bullet / List Number definitions for
    # maximum compatibility, then align them to the selected preset tokens.
    set_abstract_numbering_indent(doc, 8, left=720, hanging=360, tab_pos=720)
    set_abstract_numbering_indent(doc, 7, left=720, hanging=360, tab_pos=720)
    bullet_num_id = 1

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("TB PROJECT · INTERNAL WORKING BRIEF")
    set_run_font(hr, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    fr = fp.add_run("대한결핵협회 × ㈜마이허브  |  ")
    set_run_font(fr, size=8.5, color=MUTED)
    add_field(fp, "PAGE")

    # First-page memo masthead.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(16)
    kicker.paragraph_format.space_after = Pt(7)
    kr = kicker.add_run("PROJECT SYNTHESIS")
    set_run_font(kr, size=9.5, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    tr = title.add_run("TB 프로젝트 사업계획 종합 정리")
    set_run_font(tr, size=25, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    sr = subtitle.add_run("결핵 진단 AI 모델 공동개발의 실행 기준, 책임 체계, 인허가·기술이전 및 핵심 의사결정")
    set_run_font(sr, size=12.5, color=MUTED)
    set_paragraph_border(subtitle, color="C8D6E5", size="8", space="8")

    metadata = [
        ("기준 자료", "2026-08-26 전달 폴더 내 사업계획서 및 최신 2026-06-19 RA 반영본"),
        ("작성 기준일", "2026-08-28"),
        ("문서 용도", "내부 실행계획 정렬, 계약·인허가·데이터·개발 업무 간 기준선 확정"),
        ("해석 원칙", "사업계획서를 1차 기준으로 하고, 연구계획·계약·보안 자료는 보완 및 충돌 점검에 사용"),
    ]
    meta_tbl = doc.add_table(rows=0, cols=2)
    for label, value in metadata:
        row = meta_tbl.add_row()
        set_row_cant_split(row)
        cells = row.cells
        set_cell_text(cells[0], label, bold=True, color=DARK_BLUE, size=9.5)
        set_cell_text(cells[1], value, size=9.5)
        set_cell_shading(cells[0], BLUE_GRAY)
    apply_table_geometry(
        meta_tbl,
        [1800, 7560],
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS,
    )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)

    add_callout(
        doc,
        "핵심 결론",
        "본 과제의 성공 기준은 ‘5,000매 데이터 구축’ 자체가 아니라, 참조표준이 고정된 데이터셋, 독립 검증에서의 성능 목표 충족, 협회 환경에서의 재학습·추론 재현, 그리고 인허가 제출자료까지 이어지는 추적 가능한 개발체계를 확보하는 것이다.",
        fill=PALE_BLUE,
    )

    add_caption(doc, "표 1. 프로젝트 핵심 스냅샷")
    add_table(
        doc,
        ["구분", "문서 기준 내용"],
        [
            ("목표 제품", "흉부 X-ray 기반 결핵 선별·진단보조 독립형 디지털의료기기소프트웨어"),
            ("핵심 기능", "결핵 의심 병변 검출, 활동성/비활동성/정상 분류, Heatmap, 질환 예측 확률"),
            ("학습 데이터", "협회 수탁 영상 중 학습 가치가 높은 5,000매 선별, 비가역적 익명화 및 전문의 검수"),
            ("성능 목표", "WHO TPP Test type 1 참고: 민감도 ≥ 90%, 특이도 ≥ 80% (개발 목표이며 보장 아님)"),
            ("개발 체계", "대한결핵협회: 데이터·임상검증·내재화 / 마이허브: 툴·모델·인허가 자료·기술이전"),
            ("기간", "사업계획 2026~2028 / 계약 2028-12-31 / 연구계획 2026-06~2029-06 — 기준 통일 필요"),
            ("예산 표기", "1,000만원: 의료자문 및 학습용 영상 레이블링 인력지원비"),
        ],
        [1900, 7460],
        shade_first_col=True,
        font_size=9.1,
    )
    add_source(doc, "[BP] [RP] [CT]")

    add_heading(doc, "우선 확정이 필요한 5가지", 2)
    priorities = [
        "단일 일정 기준선: 2028년 종료인지 2029년 6월 종료인지, 연구 완료와 인허가 완료를 분리하여 확정",
        "참조표준·레이블 SOP: 활동성/비활동성/정상 정의, 확진 근거, 영상-검사 시점 허용범위, 전문의 합의 규칙 확정",
        "검수 기준: 민감도·특이도 외 AUC, 신뢰구간, 하위그룹, 재현성 허용오차를 계약 별첨으로 수치화",
        "권리 조항 정합성: ODA 등 제3자 제공 권한과 산출물 재배포 금지 조항의 충돌 해소",
        "인허가·운영 경계: 제조자/신청 주체, 협회 파인튜닝의 변경관리, UI·운영시스템 및 판독시스템 연계 범위 별도 확정",
    ]
    for item in priorities:
        add_list_paragraph(doc, item, bullet_num_id, after=5)

    doc.add_page_break()

    add_heading(doc, "1. 사업 정의와 최종 도달상태", 1)
    add_body(
        doc,
        "본 사업은 대한결핵협회가 보유한 결핵 검진 흉부 X-ray 데이터와 ㈜마이허브의 AI 개발 역량을 결합하여, 결핵 선별·진단보조에 특화된 모델을 공동 개발하고 협회가 독자적으로 재학습·검증할 수 있는 기술 기반을 확보하는 공동연구다. 상용 범용 흉부 AI의 결핵 특화 성능과 폐쇄형 운영 구조의 한계를 줄이고, 국가 결핵관리·ODA 등 공익사업에 활용 가능한 자산을 확보하는 것이 핵심 취지다.",
        align="justify",
    )

    add_heading(doc, "1.1 최종 도달상태", 2)
    final_states = [
        "제품: 흉부 X-ray에서 결핵 의심 병변과 의심 가능성을 제시하는 선별·진단보조 소프트웨어",
        "임상 출력: 활동성 결핵, 비활동성 결핵, 정상군 판별과 병변 Heatmap 및 확률 점수",
        "기술 자립: 모델 가중치, 전처리·학습·검증·추론 코드, Docker 기반 개발환경과 학습로그 확보",
        "조직 역량: 협회 내부 인력이 동일 환경에서 재학습과 추론을 재현하고 기술이전 확인서에 서명",
        "규제 준비: 식약처 신청을 위한 기술문서, 성능·위험·사용적합성·사이버보안 자료 체계 확보",
        "현장 적용: 인허가 후 협회 판독시스템과 연계하여 결핵 의심환자의 조기 선별 지원",
    ]
    for item in final_states:
        add_list_paragraph(doc, item, bullet_num_id, after=5)

    add_heading(doc, "1.2 범위와 제외 범위", 2)
    add_caption(doc, "표 2. 사업 범위 기준")
    add_table(
        doc,
        ["구분", "포함", "제외·별도 합의"],
        [
            ("데이터", "5,000매 선별, 익명화, 정밀 레이블링, 품질관리", "원목적 외 활용, 제3자 판매·양도, 무단 복제·반출"),
            ("개발", "레이블링 툴, 모델 설계·학습·고도화, SDK/API, 재학습 환경", "서비스 UI, 운영관리 시스템 등 운영 제품 개발"),
            ("기술이전", "코드·가중치·환경·로그·문서, 1회 이상 교육, 협회 환경 재현", "개발사의 범용 프레임워크·공통 모듈·선행기술"),
            ("활용", "인허가 전 연구·품질관리·교육, 인허가 후 내부·공익 활용", "인허가 전 국내 진단·치료 목적 임상 사용"),
            ("사업화", "개발사 상용화, 협회 국가사업·ODA 등 공익 활용, 독자 개량기술 사업화", "협회의 공동 산출물 직접 판매·유료 라이선싱"),
        ],
        [1300, 4030, 4030],
        aligns=["center", "left", "left"],
        font_size=8.9,
    )
    add_source(doc, "[BP] [RP] [CT]")

    add_heading(doc, "2. 추진 체계와 책임 분담", 1)
    add_body(
        doc,
        "연구계획서상 참여자는 총 9명(연구책임자 2명, 연구원 6명, 자문위원 1명)이다. 협회는 연구 총괄·데이터·의학적 검증·내부 재현을, 마이허브는 모델·도구·인허가 자료·기술이전을 주도한다.",
    )
    add_caption(doc, "표 3. 업무분장과 완료 증빙")
    add_table(
        doc,
        ["업무", "대한결핵협회", "㈜마이허브", "전문의/검수", "완료 증빙"],
        [
            ("행정·거버넌스", "질병관리청/IRB/DRB 등 검토, 연구 총괄", "계약·인허가 협력", "연구윤리·의학 자문", "승인 공문, 심의 결과, 계약"),
            ("데이터", "선별·적법한 익명화·제공, 스키마 정의", "승인 환경 수령·보안 준수", "대상·제외 기준 자문", "데이터 목록, 익명화 확인"),
            ("레이블링", "인력·업무 운영", "전용 툴 개발·이관", "기준 수립·전수 검수", "SOP, 검수기록, 라벨 버전"),
            ("모델 개발", "내부 개발자 공동 참여, 인터페이스 정의", "아키텍처·학습·최적화·고도화", "의학적 오류 분석", "코드, 가중치, 로그, 실험보고"),
            ("검증", "독립 테스트셋으로 30일 내 검수", "결함·재현성 보완", "임상적 유효성 평가", "성능보고서, 재현 확인서"),
            ("인허가", "데이터·의료진 자문 등 지원", "비용·책임 부담, 신청자료 작성·보완", "임상·사용목적 검토", "신청서, 기술문서, 보완이력"),
            ("기술이전", "협회 환경 재현·인수", "산출물 인도·1회 이상 교육", "필요 시 검토", "인수목록, 교육기록, 서명 확인서"),
        ],
        [1350, 2110, 2110, 1740, 2050],
        aligns=["center", "left", "left", "left", "left"],
        font_size=8.4,
    )
    add_source(doc, "[RP] [CT]")

    add_heading(doc, "2.1 핵심 인력", 2)
    add_body(doc, "대한결핵협회: 신우진(연구 총괄), 우성재(모델개발 참여·사업계획상 주 2~3회 개발사 파견), 김희진(의학 자문·영상 레이블링).")
    add_body(doc, "㈜마이허브: 양혁(기관 총괄), 이정훈(의학적 검토), 김상기(연구개발 총괄), 이태선(모델개발 총괄), 신동훈(모델·레이블링 툴), 최병윤(모델개발).")
    add_callout(doc, "확정 필요", "계약서는 협회 내부 개발인력 2인의 참여를 요구하지만 연구분담표에는 모델개발 참여자가 우성재 1인으로 명시되어 있다. 두 번째 인력과 투입률, 파견 방식, 인수 책임을 확정해야 한다.", fill=PALE_GOLD, label_color=GOLD)

    add_heading(doc, "3. 단계별 로드맵과 게이트", 1)
    add_caption(doc, "표 4. 사업계획서 기준 단계별 로드맵")
    add_table(
        doc,
        ["단계", "기간", "핵심 작업", "단계 종료 기준"],
        [
            ("1. 행정 승인·인프라", "8주", "질병관리청 승인, 공동연구계약, 전문의 위촉, 레이블링 툴 개발·이관", "필수 승인·계약 완료, 승인 환경과 툴 사용 가능"),
            ("2. 데이터 선별·익명화", "4주", "학습 가치가 높은 영상 5,000매 선별, 비가역적 익명화", "대상목록·익명화 검증·반출 승인 완료"),
            ("3. 정밀 레이블링", "20주", "병변 위치·특징 정의, 활동성/비활동성/정상 분류, 전문의 전수 검수", "레이블 잠금, 검수기록, 품질지표 승인"),
            ("4. AI 학습·내재화", "20주", "파견 공동개발, Fine-tuning, 하이퍼파라미터 최적화, 반복 성능개선", "모델 동결, 실험로그, 독립검증 준비"),
            ("5. 이관·인허가", "2027-05~2028", "1차 산출물 이관·재현 검증, 인허가 병행, 승인 반영 최종 이관", "연구 완료 확인서 / 승인 후 10일 내 최종 이관"),
        ],
        [1500, 1150, 4090, 2620],
        aligns=["center", "center", "left", "left"],
        font_size=8.8,
    )
    add_source(doc, "[BP] 단계 1~5, [CT] 검수·이관 조항")

    add_heading(doc, "3.1 권고 마일스톤", 2)
    milestones = [
        "M0 착수 승인: 계약, 연구심의, 데이터 처리·보안 계획, 제품 사용목적 기준선 승인",
        "M1 레이블 파일럿 승인: 소량 파일럿을 통해 라벨 정의·툴·전문의 일치도 검증 (수량은 별도 확정)",
        "M2 데이터셋 잠금: 학습·검증·시험 분리와 중복·누수 점검 완료",
        "M3 모델 동결: 버전·코드·가중치·환경·학습로그를 하나의 릴리스로 고정",
        "M4 연구 완료: 독립 성능검증, 협회 환경 재현, 산출물 이관·교육·확인서 완료",
        "M5 인허가 제출: 신청 주체의 QMS/GMP와 기술문서·시험보고서 패키지 제출",
        "M6 승인·최종 이관: 승인사항 반영 모델과 문서를 승인 후 10일 내 재이관",
    ]
    milestone_num_id = add_numbering_instance(doc, 7)
    for item in milestones:
        add_list_paragraph(doc, item, milestone_num_id, after=5)

    add_callout(
        doc,
        "일정 불일치",
        "사업계획은 2026~2028, 계약은 2028-12-31, 연구계획은 2026-06~2029-06으로 서로 다르다. 연구개발 완료일(M4)과 인허가 완료 목표일(M6)을 분리하고, 계약기간·데이터 처리기간·무상지원 기산점을 하나의 일정표에 연결해야 한다.",
        fill=PALE_RED,
        label_color=RED,
    )

    add_heading(doc, "4. 데이터·모델·성능 검증 설계", 1)
    add_heading(doc, "4.1 데이터셋 구성", 2)
    add_caption(doc, "표 5. 데이터 구축 기준과 확정 과제")
    add_table(
        doc,
        ["항목", "문서 기준", "착수 전 확정할 사항"],
        [
            ("규모", "학습 가치가 높은 5,000매", "클래스별 수량, 기관·장비·연령 분포, 추가 데이터 조건"),
            ("분류", "활동성 결핵, 비활동성 결핵, 정상, 필요 시 기타 폐질환", "상호배타성, 다중병변·복합질환 처리 규칙"),
            ("정답 라벨", "병변 위치·유형, 연령대(10세 단위), 전문의 최종 검수", "확진검사·판독·추적결과 우선순위와 합의 절차"),
            ("품질관리", "오진·중복·저품질 영상 제거", "제외 사유 코드, 중복 기준, 재검수·감사 샘플링"),
            ("분할", "학습과 독립된 검증 데이터셋", "환자 단위 분리, 시험셋 봉인, 기관·장비 외부검증"),
            ("익명화", "성명·생년월일·기관명 등 식별정보 비가역 제거", "DICOM 태그·픽셀 번인·연결키·검증 책임자 체크리스트"),
        ],
        [1450, 3580, 4330],
        aligns=["center", "left", "left"],
        font_size=8.8,
    )
    add_source(doc, "[BP] [RP] [DS] [MN]")

    add_heading(doc, "4.2 모델 기능과 성능 기준", 2)
    add_caption(doc, "표 6. 최소 성능·검증 항목")
    add_table(
        doc,
        ["구분", "목표/요건", "검증 포인트"],
        [
            ("민감도", "≥ 90%", "독립 시험셋, 신뢰구간, 활동성 결핵 우선 분석"),
            ("특이도", "≥ 80%", "정상군 및 비결핵 폐질환에서의 오탐 분석"),
            ("AUC", "평가 필수, 수치 미정", "ROC/AUC 및 임계값 선정 근거 확정"),
            ("오류분석", "위양성·위음성 집중 분석", "장비·기관·선량·연령·영상 품질 하위그룹"),
            ("출력 기능", "Heatmap, 질환 예측 확률", "표시 방식, API 필드, 실패·불확실 결과 처리"),
            ("재현성", "협회 환경에서 학습·추론 재현", "허용오차, seed·버전·하드웨어 차이 처리"),
            ("인허가", "신청 1건", "성능평가·임상적 유효성·변경관리 자료 연계"),
        ],
        [1500, 3270, 4590],
        aligns=["center", "left", "left"],
        font_size=8.8,
    )
    add_source(doc, "[RP] WHO TPP 참고 목표. 목표치는 최종 성능 보장을 의미하지 않음.")

    add_callout(
        doc,
        "검수 기준 권고",
        "계약은 ‘식약처 가이드라인의 성능 목표 충족’을 요구하지만 구체 수치와 통계 기준이 없다. 민감도 90%·특이도 80%를 수용한다면 시험셋 정의, 분석단위, 신뢰구간 하한, 임계값 고정 시점, 결측·품질불량 처리와 재시험 조건까지 별첨으로 명시해야 한다.",
        fill=PALE_GOLD,
        label_color=GOLD,
    )

    doc.add_page_break()

    add_heading(doc, "5. 식약처 인허가 추진체계", 1)
    add_body(
        doc,
        "최신 RA 반영본은 본 모델을 결핵 확진기가 아니라 의료인의 결핵 선별·진단 의사결정을 보조하는 독립형 디지털의료기기소프트웨어로 정의한다. 예상 제품코드는 B1ABXA1로 검토하되, 최종 제품코드·등급·허가/인증 대상·제출자료 범위는 식약처 또는 심사기관의 판단에 따른다.",
        align="justify",
    )
    add_caption(doc, "표 7. 인허가 단계와 준비물")
    add_table(
        doc,
        ["단계", "주요 활동", "핵심 산출물"],
        [
            ("1. 사용목적·품목분류", "입력·출력·사용자·환경·임상흐름 정의, 필요 시 사전상담", "제품정의서, 품목분류 검토, 상담 기록"),
            ("2. 설계·개발 문서", "알고리즘·모델·소프트웨어 구조, 요구사항·설계·형상·변경관리", "제품개요, SRS, SDS, 개발·변경관리 문서"),
            ("3. 데이터·성능", "수집·선정·분리·익명화·레이블·참조표준과 통계계획 관리", "데이터 문서, 성능평가 계획·보고, 임상 유효성 자료"),
            ("4. 안전·품질", "위험관리, 사용적합성, 사이버보안·개인정보, QMS/GMP 준비", "위험관리 파일, 시험·검토보고서, QMS/GMP 자료"),
            ("5. 신청·보완", "신청서 제출, 질의·보완, 승인사항 반영", "신청·보완 이력, 사용자매뉴얼, 표시·주의사항"),
        ],
        [1650, 4440, 3270],
        aligns=["center", "left", "left"],
        font_size=8.8,
    )
    add_source(doc, "[RP] 4-6, [CT] 제9조, [RA] 인허가 관련 사항 정리")

    add_heading(doc, "5.1 책임·제약", 2)
    for item in [
        "비용과 책임: 인허가 및 신의료기술평가 등 제반 절차는 마이허브가 부담·수행하고, 협회는 데이터와 의료진 자문을 지원한다.",
        "명의: 개발사가 취득한 의료기기 품목 허가 및 인증 명의는 개발사에 귀속된다.",
        "연구 완료: 인허가 획득은 연구 완료 판정의 필수조건이 아니며, 성능·재현성·이관 조건 충족 시 먼저 완료 가능하다.",
        "임상 사용: 인허가 전에는 국내 의료현장의 진단·치료 목적으로 사용할 수 없고, 연구·고도화·품질관리·교육에 한해 사용한다.",
        "최종 이관: 인허가 취득 후 10일 이내 승인사항이 반영된 모델과 문서를 다시 이관하며, 기능상 오류에 대해 1년 무상지원한다.",
    ]:
        add_list_paragraph(doc, item, bullet_num_id, after=5)

    add_callout(
        doc,
        "규제 운영 리스크",
        "협회가 승인 이후 자체 파인튜닝·재학습을 수행하면 허가받은 모델과 연구용 모델이 달라질 수 있다. 제조자 책임, 변경관리계획, 릴리스 승인권, 현장 배포 가능한 버전의 식별 규칙을 인허가 전략과 함께 설계해야 한다.",
        fill=PALE_RED,
        label_color=RED,
    )

    add_heading(doc, "6. 산출물·검수·기술이전", 1)
    add_caption(doc, "표 8. 이관 산출물 체크리스트")
    add_table(
        doc,
        ["산출물", "포함 범위", "인수 확인"],
        [
            ("소스코드", "전처리, 학습, 검증, 추론, Probability Score, Heatmap", "저장소 태그·커밋, 빌드·실행 절차"),
            ("모델 가중치", "사전학습 모델, 최종 파라미터, 체크포인트", "해시, 버전, 성능보고서 연결"),
            ("개발 환경", "재학습 가능한 Docker 이미지와 환경 설정", "협회 인프라 설치·재현 결과"),
            ("학습 로그", "곡선, 손실·정확도, 하이퍼파라미터, 실험 기록", "최종 릴리스와 로그 일치"),
            ("기술 문서", "최종보고서, API 명세서, 설치·사용자 매뉴얼", "문서 버전·승인자·변경이력"),
            ("실행 인터페이스", "AI 모델 SDK/API", "입출력 스키마, 오류코드, 성능·보안 조건"),
            ("레이블링 툴", "실행파일 또는 사용권한", "사용조건·버전·지원범위; 권리는 개발사 귀속"),
        ],
        [1800, 4580, 2980],
        aligns=["center", "left", "left"],
        font_size=8.8,
    )
    add_source(doc, "[CT] 제4조·제10조 및 별첨1")

    add_heading(doc, "6.1 검수 순서", 2)
    acceptance_num_id = add_numbering_instance(doc, 7)
    for item in [
        "개발사가 목표 기능과 산출물을 제출하고 검수를 요청한다.",
        "협회는 제출일부터 30일 이내 별도 테스트셋으로 성능을 검증한다.",
        "협회 내부 개발자가 제공 환경에서 재학습·추론을 수행해 재현성을 확인한다.",
        "미달·비재현 시 개발사가 즉시 수정·보완하고 재검수를 받는다.",
        "성능, 재현성, 산출물 이관, 교육이 모두 완료되면 양측이 이관·재현 확인서에 서명한다.",
    ]:
        add_list_paragraph(doc, item, acceptance_num_id, after=5)

    add_callout(
        doc,
        "코드 경계 확정",
        "연구계획서는 개발사의 범용 프레임워크·라이브러리·공통 모듈·선행기술을 이전 대상에서 제외한다. 반면 계약은 추론 전 과정의 소스코드 원본을 요구한다. 재현에 필요한 코드·컨테이너는 누락되지 않도록 ‘Background IP 목록’과 대체 배포방식(바이너리/라이선스)을 별첨해야 한다.",
        fill=PALE_GOLD,
        label_color=GOLD,
    )

    add_heading(doc, "7. 데이터 보안과 생명주기", 1)
    add_caption(doc, "표 9. 데이터 통제 항목")
    add_table(
        doc,
        ["통제영역", "문서상 요구사항", "보완 필요"],
        [
            ("접근", "최소인원, 데이터관리담당자 지정, 승인 IP만 접근", "역할별 권한표, 승인·회수 절차, 접근로그 보존기간"),
            ("저장", "승인된 저장공간·환경, 개인 PC/클라우드/이메일 금지", "암호화 기준, 키관리, 백업·복구, 개발·시험 환경 분리"),
            ("반출", "무단 복제·외부매체·제3자 제공 금지", "승인 워크플로, 파일 해시, 반출대장, 원격접속 통제"),
            ("물리", "승인자만 서버·보관장소 접근, 출력 최소화·표기", "출입기록, 화면·촬영 통제, 출력물 회수 확인"),
            ("사고", "유출·오남용 의심 시 즉시 보고, 중단·원인·재발방지", "연락망, 초기보고 시간, 포렌식·대외 신고 기준"),
            ("점검", "협회가 연 1회 이상 이행 여부 점검", "정기 증빙 목록, 시정조치 기한, 협력사 재점검"),
            ("파기", "연구종료·계약해지 후 5일 내 복구불가 파기 및 확약서", "연구완료/계약종료 트리거 통일, 파기 검증방법"),
            ("규제 보관", "법적 필수 데이터는 서면 통보·망분리 암호화·기간 후 파기", "보존항목·법적 근거·기간·책임자 목록"),
        ],
        [1450, 4050, 3860],
        aligns=["center", "left", "left"],
        font_size=8.6,
    )
    add_source(doc, "[CT] 제5~6조, [DS]")

    add_heading(doc, "8. 지식재산·활용·사업화 구조", 1)
    add_caption(doc, "표 10. 권리 구조 요약")
    add_table(
        doc,
        ["대상", "권리 귀속", "활용 범위·제한"],
        [
            ("원천·가공·레이블·파생 데이터", "대한결핵협회", "개발사는 계약기간 중 과제 목적의 한시적 이용만 가능"),
            ("레이블링 툴", "㈜마이허브", "협회에는 실행파일 또는 사용권한 부여"),
            ("공동 연구 산출물", "협회·개발사 50:50 공동소유", "개발사는 독자 상용화·수익 귀속; 협회는 비영리·공익 활용"),
            ("협회의 공익사업", "협회 사용권", "내부 판독, 국가사업, ODA·국제협력에 제공·재실시권 부여 가능"),
            ("독자 개량기술", "개발한 당사자 단독소유", "목적 제한 없이 상용화·라이선싱 가능"),
            ("인허가 명의", "㈜마이허브", "인허가권으로 협회의 계약상 내부·공익 사용을 방해할 수 없음"),
        ],
        [2050, 2050, 5260],
        aligns=["center", "center", "left"],
        font_size=8.8,
    )
    add_source(doc, "[CT] 제5조·제7~9조")

    add_heading(doc, "8.1 계약 정합성 점검", 2)
    for item in [
        "제7조는 국가사업·ODA를 위해 제3자에게 산출물을 제공하고 재실시권을 부여할 수 있도록 하나, 제10조는 별도 합의 없는 제3자 제공·재배포를 금지한다.",
        "공동소유 50:50임에도 공동 산출물의 상용화 권한과 수익이 개발사에 집중된다. 지분의 의미, 의사결정권, 특허비용과 유지의무를 별도로 정해야 한다.",
        "‘독자적 비용과 노력’으로 만든 개량기술의 구분 기준이 없어 공동 인력·공동 데이터·공동 인프라를 사용한 경우 분쟁 소지가 있다.",
        "협회 자체 재학습 모델의 상용화 권리는 인정되지만, 공동 산출물·데이터·개발사 선행기술 사용 범위와 인허가 변경관리 의무가 연결되어 있지 않다.",
    ]:
        add_list_paragraph(doc, item, bullet_num_id, after=5)

    doc.add_page_break()

    add_heading(doc, "9. 예산·자원 계획", 1)
    add_caption(doc, "표 11. 문서상 확인되는 비용·자원")
    add_table(
        doc,
        ["항목", "문서상 기준", "확정 필요"],
        [
            ("연구비", "총 10,000천원(1,000만원), 의료자문 및 학습용 영상 레이블링 인력지원비 100%", "지급조건, 검수기준, 세금·부대비용, 추가 데이터 비용"),
            ("인허가", "마이허브의 비용과 책임", "임상·시험·GMP·사이버보안·컨설팅 비용 범위와 상한"),
            ("협회 인력", "총괄·실무·자문, 내부 개발인력 2인 공동 참여", "두 번째 개발인력, 파견 2~3회/주 기간, 인프라 운영담당"),
            ("개발사 인력", "기관·R&D·모델·툴·의학검토 담당", "단계별 투입계획, PM/QA/RA/보안 담당자"),
            ("인프라", "협회 재학습 환경 구축", "GPU·저장소·백업·망분리·MLOps·유지비"),
            ("현장 연계", "협회 판독시스템 적용 목표", "UI·운영관리 제외에 따른 별도 연계 개발·시험·운영 예산"),
        ],
        [1700, 3850, 3810],
        aligns=["center", "left", "left"],
        font_size=8.8,
    )
    add_source(doc, "[BP] [RP] [CT]")

    add_callout(
        doc,
        "예산 해석",
        "현재 명시된 1,000만원은 전체 AI 개발비라기보다 전문의 자문·레이블링 인력지원비로 보인다. 데이터·컴퓨팅·통합·규제시험·QMS/GMP·보안·운영 전환 비용을 포함한 총사업비와 현물투입을 별도 WBS 예산으로 작성해야 한다.",
        fill=PALE_GOLD,
        label_color=GOLD,
    )

    add_heading(doc, "10. 핵심 리스크와 의사결정 목록", 1)
    add_caption(doc, "표 12. 우선순위 리스크·결정 매트릭스")
    add_table(
        doc,
        ["우선", "이슈", "영향", "필요 결정", "완료 게이트"],
        [
            ("상", "연구기간 불일치", "계약·데이터 처리·지원기간 충돌", "M4 연구완료와 M6 인허가 완료 일자 분리", "계약 확정 전"),
            ("상", "활동성/비활동성 참조표준", "라벨 오류·성능 왜곡", "확진 근거·시점·전문의 합의 규칙 승인", "레이블링 전"),
            ("상", "시험셋·성능 검수 정의 미흡", "완료판정 분쟁", "수치·통계·하위그룹·재시험 조건 별첨", "학습 전"),
            ("상", "IP/ODA 조항 충돌", "공익사업 배포 제약·분쟁", "제7조와 제10조 우선순위·예외 통일", "계약 확정 전"),
            ("상", "인허가 모델과 재학습 모델 혼재", "무허가 변경·책임 혼선", "릴리스·변경관리·배포 승인권 정의", "설계 동결 전"),
            ("상", "익명화·규제보관 트리거 불명확", "개인정보·계약 위반", "파기·보존 대상과 기산일·증빙 통일", "데이터 반출 전"),
            ("중", "5,000매 대표성·불균형", "일반화 성능 저하", "클래스·기관·장비·연령 샘플링 계획", "데이터 잠금 전"),
            ("중", "UI·운영시스템 범위 제외", "현장 적용 지연", "판독 연계/API/운영 프로젝트와 예산 분리", "MVP 설계 전"),
            ("중", "선행기술·소스코드 경계", "재현 실패·라이선스 분쟁", "Background IP 및 대체 제공 방식 별첨", "개발 착수 전"),
            ("중", "인력 2인 요건 불명확", "기술내재화 미흡", "두 번째 인력과 투입·교육·인수 책임 지정", "착수 전"),
            ("중", "총사업비 미정", "규제·인프라 중단", "단계별 현금·현물·외부시험 예산 승인", "기준선 승인 전"),
            ("중", "장비·저선량·기관 편차", "현장 성능 변동", "하위그룹 검증·외부검증·모니터링 계획", "검증계획 승인 전"),
        ],
        [700, 2000, 2260, 2850, 1550],
        aligns=["center", "left", "left", "left", "center"],
        font_size=8.2,
    )
    add_source(doc, "[BP] [RP] [CT] [DS] [MN]")

    add_heading(doc, "11. 30일 실행안", 1)
    add_body(doc, "아래는 문서 정합성과 착수 준비를 위한 권고 실행안이다. 실제 착수 상태를 의미하지 않으며, 담당자·기한은 양 기관이 확정해야 한다.", italic=True, color=MUTED)

    add_heading(doc, "1주차 — 기준선 확정", 2)
    for item in [
        "사업계획·연구계획·계약의 기준 버전과 우선순위를 지정하고 변경대장을 개설한다.",
        "연구 완료일·인허가 목표일·데이터 처리기간·기술지원 기산점을 통합 일정표에 고정한다.",
        "제품 사용목적, 대상 사용자·환자군, 입력·출력, 연구용/허가용 모델 구분을 1페이지 제품정의서로 승인한다.",
    ]:
        add_list_paragraph(doc, item, bullet_num_id, after=4)

    add_heading(doc, "2주차 — 데이터·레이블 체계", 2)
    for item in [
        "데이터 사전, DICOM 익명화 체크리스트, 접근·반출·파기대장과 데이터관리담당자를 확정한다.",
        "활동성/비활동성/정상 참조표준과 확진검사·시점·전문의 불일치 처리 SOP를 승인한다.",
        "5,000매 샘플링·분할 계획과 레이블 파일럿 승인 기준을 수립한다.",
    ]:
        add_list_paragraph(doc, item, bullet_num_id, after=4)

    add_heading(doc, "3주차 — 기술·검증·이관 설계", 2)
    for item in [
        "모델·데이터 파이프라인, SDK/API, 판독시스템 Interface Table과 운영 범위 경계를 설계한다.",
        "성능평가·통계계획, 독립 시험셋 봉인, 하위그룹 분석, 재현성 허용오차를 확정한다.",
        "산출물 구조, Background IP, 라이선스, Docker 재현 절차와 인수 체크리스트를 작성한다.",
    ]:
        add_list_paragraph(doc, item, bullet_num_id, after=4)

    add_heading(doc, "4주차 — 계약·인허가·예산 승인", 2)
    for item in [
        "ODA/제3자 제공, 공동소유·상용화, 개량기술, 소스코드 경계, 데이터 보존·파기 충돌을 계약 별첨에 반영한다.",
        "식약처 사전상담 질문서와 제출자료 갭 분석을 완료하고 신청 주체·QMS/GMP·변경관리 책임을 확정한다.",
        "개발·레이블·컴퓨팅·보안·통합·규제시험·운영전환을 포함한 WBS 예산과 의사결정 게이트를 승인한다.",
    ]:
        add_list_paragraph(doc, item, bullet_num_id, after=4)

    add_callout(
        doc,
        "30일 종료 조건",
        "단일 기준 일정, 승인된 제품정의서, 데이터·레이블 SOP, 성능검수 별첨, 계약 쟁점 수정안, 인허가 갭 분석, WBS 예산이 모두 문서화되어야 본 개발 착수 기준선이 성립한다.",
        fill=PALE_GREEN,
        label_color=GREEN,
    )

    add_heading(doc, "부록 A. 기준 문서와 사용 원칙", 1)
    add_caption(doc, "표 A-1. 본 정리본에 사용한 자료")
    add_table(
        doc,
        ["표기", "문서", "사용 범위"],
        [
            ("[BP]", "결핵진단 AI 사업계획서.hwp", "5단계 로드맵, 5,000매, 전문의 예산, 이관·인허가 기본 원칙"),
            ("[RP]", "대한결핵협회 연구과제 신청서 및 계획서(2026.6.19)_RA.hwpx", "목표·방법·성능·인허가·연구기간·예산·인력의 최신 RA 반영안"),
            ("[CT]", "공동연구개발계약서(안)_마이허브_대한결핵협회(2026.6.19)_RA.hwp", "역할·데이터·IP·상용화·검수·이관·지원·책임"),
            ("[DS]", "연구 데이터의 안전성 확보 조치 및 파기 계획서.docx", "접근·저장·사고·점검·파기 통제"),
            ("[RA]", "인허가 관련 사항 정리.hwpx", "인허가 처리·제출자료 체크 항목"),
            ("[MN]", "참고_대한결핵협회 CXR TB 모델관련 회의_20251211.docx", "라벨·데이터 편차·변경관리·사업화의 초기 리스크 배경"),
        ],
        [900, 4880, 3580],
        aligns=["center", "left", "left"],
        font_size=8.6,
    )

    add_heading(doc, "자료 적용 원칙", 2)
    for item in [
        "동일한 사업계획서 2개 사본은 SHA-256이 동일하여 한 번만 분석했다.",
        "인허가 관련 표현은 2026-06-19 ‘RA’ 반영본을 최신안으로 사용했다.",
        "회의록은 확정 지시가 아니라 과거 논의와 리스크 식별 자료로만 사용했다.",
        "법령·품목분류·심사자료는 실제 신청 전 RA/법무 및 관계기관 확인이 필요하다.",
        "‘권고’, ‘확정 필요’, ‘보완 필요’ 표기는 원문 확정사항이 아니라 문서 간 대조를 통해 도출한 실행 제안이다.",
    ]:
        add_list_paragraph(doc, item, bullet_num_id, after=5)

    add_source(doc, "본 문서는 2026-08-26 제공자료의 내부 정리본이며, 계약·규제·의학적 판단을 대체하지 않는다.")

    doc.save(OUTPUT)
    issues = audit_docx_tables(OUTPUT)
    if issues:
        raise SystemExit(f"Table geometry audit failed with {issues} issue(s)")
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
